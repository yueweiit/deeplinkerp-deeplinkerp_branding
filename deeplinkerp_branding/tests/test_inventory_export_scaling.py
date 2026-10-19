from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from deeplinkerp_branding.services import inventory_detail_service as service


@pytest.mark.parametrize("categorized", [False, True])
def test_export_formats_fixed_columns_without_scanning_all_existing_cells(monkeypatch, categorized):
	monkeypatch.setattr(service, "frappe", None)
	stock = [{"item_code": "I", "warehouse": "W", "actual_qty": 5, "stock_uom": "kg"}]
	snapshots = [
		{"item_code": "I", "warehouse": "W", "original_location": location, "location_qty": qty}
		for location, qty in [("L1", 2), ("L2", 3)]
	]
	if categorized:
		payload = service.build_categorized_inventory_payload(stock, snapshots, category="semi_finished")
		build = service.build_categorized_inventory_xlsx
	else:
		payload = service.build_inventory_location_payload(snapshots, {}, page_length=None)
		build = service.build_inventory_location_xlsx
	reads = 0
	original = Worksheet.max_column.fget

	def max_column(sheet):
		nonlocal reads
		reads += 1
		return original(sheet)

	with patch.object(Worksheet, "max_column", property(max_column)):
		content = build(payload)
	# Serialization may inspect dimensions, but formatting must not do it per row.
	assert reads <= 1
	book = load_workbook(BytesIO(content))
	sheet = book.active
	assert sheet.max_row == 3
	assert sheet.max_column == (14 if categorized else 11)
	assert sheet["F2"].value == 5
	assert sheet["E2"].value + sheet["E3"].value == 5
	assert "A2:A3" in sheet.merged_cells
	assert sheet["A2"].alignment.wrap_text
	assert sheet["E3"].alignment.vertical == "center"
