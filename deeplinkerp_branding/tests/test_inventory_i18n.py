"""Exercise localized inventory rendering and exports with the shipped catalogs."""

from __future__ import annotations

import ast
import csv
import json
import re
import subprocess
from io import BytesIO
from pathlib import Path
from string import Formatter
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from deeplinkerp_branding.services import inventory_detail_service as service

PACKAGE = Path(__file__).resolve().parents[1]


def catalog(language):
	with (PACKAGE / "translations" / f"{language}.csv").open(newline="", encoding="utf-8") as handle:
		return {row[0]: row[1] for row in csv.reader(handle) if row and (len(row) < 3 or not row[2])}


@pytest.mark.parametrize("language", ["en", "es", "zh"])
def test_inventory_catalogs_cover_the_shared_page_and_export(language):
	translations = catalog(language)
	source = (PACKAGE / "public/js/inventory_detail.bundle.js").read_text()
	keys = set(re.findall(r'__\("([^"\n]+)"', source))
	for node in ast.walk(ast.parse(Path(service.__file__).read_text())):
		if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "_":
			if node.args and isinstance(node.args[0], ast.Constant) and isinstance(node.args[0].value, str):
				keys.add(node.args[0].value)
	keys.add("该类型需要在标准物料移动表单中补充工单、BOM、批次、序列号或物料角色等字段。")
	keys.update(service.EXPORT_HEADERS + service.CATEGORY_EXPORT_HEADERS)
	keys.update(definition["title"] for definition in service.CATEGORY_DEFINITIONS.values())
	keys.update(["物料库存明细", "库存差异", "库存一致", "库位待维护", "待库存重算", "重算中", "校验中", "冲销完成", "失败待处理"])
	for key in keys:
		assert key in translations, (language, key)
		original = {(field, spec) for _, field, spec, _ in Formatter().parse(key) if field is not None}
		localized = {(field, spec) for _, field, spec, _ in Formatter().parse(translations[key]) if field is not None}
		assert original == localized


@pytest.mark.parametrize("language", ["en", "es", "zh"])
def test_client_translates_display_text_and_preserves_business_values(language):
	translations = catalog(language)
	body = r"""
const assert = require('node:assert/strict');
const translations = JSON.parse(require('node:fs').readFileSync(0, 'utf8'));
global.__ = (message, values = []) => (translations[message] || message)
  .replace(/\{(\d+)\}/g, (token, i) => values[i] ?? token);
const formats = [];
global.frappe = {
  format(value, field, options) { formats.push({value, field, options}); return String(value); },
  datetime: {str_to_user: value => 'localized:' + value},
};
const inventory = require(SCRIPT);
const group = {item_code:'I<1>',item_name:'库存差异',warehouse:'中文仓库',actual_qty:5,
  stock_uom:'个',inventory_status:'库存差异',locations:[{reference_location:'库位待维护',snapshot_date:'2026-10-19'}]};
const html = inventory.renderCategorizedRows([group]);
assert.ok(html.includes(inventory.escapeHtml(translations['库存差异'])));
assert.ok(html.includes(translations['库位待维护']));
assert.ok(html.includes('is-difference'));
assert.ok(html.includes('中文仓库'));
assert.ok(html.includes('>库存差异</td>'));
assert.ok(html.includes('I&lt;1&gt;'));
assert.ok(html.includes('localized:2026-10-19'));
assert.equal(group.inventory_status, '库存差异');
assert.ok(formats.every(format => format.options.only_value));
assert.equal(formats[0].field.precision, 0);
inventory.formatQuantity(1.25, 'kg');
assert.equal(formats.at(-1).field.precision, 2);
assert.throws(() => inventory.normalizePageLength(0),
  error => error.message === __('每页行数必须是 {0} 到 {1} 的整数。', [1, 2500]));
const proto = inventory.InventoryDetailPage.prototype;
for (const name of ['makeFilters', 'renderShell', 'bindEvents', 'refresh']) proto[name] = () => {};
frappe.ui = {make_app_page: config => config};
const page = new inventory.InventoryDetailPage({}, {category:'finished_goods', title:'成品库存明细'});
assert.equal(page.page.title, translations['成品库存明细']);
const button = {length:1,text(value){this.label=value;},prop(){},attr(){}};
page.$movementButton = button;
page.updateMovementButton();
assert.equal(button.label, __('物料移动 ({0})', [0]));
""".replace(
		"SCRIPT", json.dumps(str(PACKAGE / "public/js/inventory_detail.bundle.js"))
	)
	result = subprocess.run(["node", "-e", body], input=json.dumps(translations), text=True, capture_output=True)
	assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("language", ["en", "es", "zh"])
def test_exports_translate_labels_and_keep_numbers_and_business_names(monkeypatch, language):
	translations = catalog(language)
	monkeypatch.setattr(service, "frappe", SimpleNamespace(_=lambda key: translations.get(key, key)))
	group = {
		"item_code": "I-1", "item_name": "库存差异", "warehouse": "中文仓库", "actual_qty": 5,
		"difference_qty": 2, "stock_uom": "个", "item_group": "业务分类", "inventory_status": "库存差异",
		"locations": [{"reference_location": "库位待维护", "snapshot_location_qty": 3}],
	}
	sheet = load_workbook(BytesIO(service.build_categorized_inventory_xlsx({"title": "成品库存明细", "groups": [group]}))).active
	assert sheet.title == translations["成品库存明细"][:31]
	assert [cell.value for cell in sheet[1]] == [translations[key] for key in service.CATEGORY_EXPORT_HEADERS]
	assert sheet["B2"].value == "库存差异"
	assert sheet["C2"].value == "中文仓库"
	assert sheet["D2"].value == translations["库位待维护"]
	assert sheet["F2"].value == 5
	assert sheet["F2"].data_type == "n"
	assert sheet["H2"].value == translations["库存差异"]
	assert sheet["K2"].value == "业务分类"
	material = load_workbook(BytesIO(service.build_inventory_location_xlsx({"groups": []}))).active
	assert material.title == translations["物料库存明细"][:31]
	assert [cell.value for cell in material[1]] == [translations[key] for key in service.EXPORT_HEADERS]
	with pytest.raises(ValueError) as error:
		service.get_category_definition("invalid")
	assert str(error.value) == translations["不支持的库存分类：{0}"].format("invalid")
